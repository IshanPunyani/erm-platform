import sqlite3, json, os
from datetime import datetime

DB_PATH = os.environ.get("DB_PATH", "erm.db")

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA foreign_keys=ON")
    return conn

def init_db():
    conn = get_db()
    c = conn.cursor()

    c.executescript("""
    CREATE TABLE IF NOT EXISTS settings (
        key   TEXT PRIMARY KEY,
        value TEXT
    );

    CREATE TABLE IF NOT EXISTS businesses (
        id         INTEGER PRIMARY KEY AUTOINCREMENT,
        name       TEXT NOT NULL UNIQUE,
        active     INTEGER DEFAULT 1,
        created_at TEXT DEFAULT (datetime('now'))
    );

    CREATE TABLE IF NOT EXISTS risks (
        id                    INTEGER PRIMARY KEY AUTOINCREMENT,
        no                    TEXT NOT NULL,
        category              TEXT,
        statement             TEXT,
        events                TEXT,
        impact_financial      TEXT,
        impact_financial_basis TEXT,
        impact_operational    TEXT,
        impact_operational_basis TEXT,
        impact_regulatory     TEXT,
        impact_regulatory_basis TEXT,
        impact_brand          TEXT,
        impact_brand_basis    TEXT,
        final_impact          INTEGER DEFAULT 0,
        likelihood            INTEGER,
        likelihood_basis      TEXT,
        risk_score            INTEGER DEFAULT 0,
        rating                TEXT DEFAULT 'N/A',
        applicability         TEXT DEFAULT 'group',
        businesses            TEXT DEFAULT '[]',
        risk_owner            TEXT,
        strategy              TEXT,
        existing_controls     TEXT,
        proposed_plan         TEXT,
        proposed_plan_submitted_by TEXT,
        proposed_plan_submitted_at TEXT,
        mitigation_owner      TEXT,
        target_date           TEXT,
        status                TEXT DEFAULT 'Not Started',
        actual_date           TEXT,
        assessment_frequency  TEXT DEFAULT 'Quarterly',
        next_assessment_date  TEXT,
        added_by              TEXT,
        created_at            TEXT DEFAULT (datetime('now')),
        updated_at            TEXT DEFAULT (datetime('now'))
    );

    CREATE TABLE IF NOT EXISTS risk_history (
        id         INTEGER PRIMARY KEY AUTOINCREMENT,
        risk_id    INTEGER NOT NULL,
        date       TEXT NOT NULL,
        impact     INTEGER,
        likelihood INTEGER,
        score      INTEGER,
        rating     TEXT,
        note       TEXT,
        changed_by TEXT,
        created_at TEXT DEFAULT (datetime('now')),
        FOREIGN KEY (risk_id) REFERENCES risks(id) ON DELETE CASCADE
    );

    CREATE TABLE IF NOT EXISTS approvals (
        id          INTEGER PRIMARY KEY AUTOINCREMENT,
        risk_id     INTEGER NOT NULL,
        status      TEXT NOT NULL,
        approver    TEXT NOT NULL,
        approved_by TEXT,
        approved_at TEXT,
        comments    TEXT,
        created_at  TEXT DEFAULT (datetime('now')),
        FOREIGN KEY (risk_id) REFERENCES risks(id) ON DELETE CASCADE
    );

    CREATE TABLE IF NOT EXISTS velocity_scale (
        id          INTEGER PRIMARY KEY AUTOINCREMENT,
        score       INTEGER NOT NULL UNIQUE,
        label       TEXT NOT NULL,
        description TEXT
    );

    CREATE TABLE IF NOT EXISTS review_periods (
        id                 INTEGER PRIMARY KEY AUTOINCREMENT,
        risk_id            INTEGER NOT NULL,
        period_label       TEXT NOT NULL,
        due_date           TEXT,
        status             TEXT NOT NULL DEFAULT 'Due',
        submitted_by       TEXT,
        submitted_at       TEXT,
        final_impact       INTEGER,
        impact_financial   TEXT,
        impact_operational TEXT,
        impact_regulatory  TEXT,
        impact_brand       TEXT,
        likelihood         INTEGER,
        velocity           INTEGER,
        score_2d           INTEGER,
        rating_2d          TEXT,
        score_3d           INTEGER,
        rating_3d          TEXT,
        comments           TEXT,
        attachments        TEXT DEFAULT '[]',
        approver_role      TEXT,
        approved_by        TEXT,
        decided_at         TEXT,
        decision_comments  TEXT,
        created_at         TEXT DEFAULT (datetime('now')),
        FOREIGN KEY (risk_id) REFERENCES risks(id) ON DELETE CASCADE
    );

    CREATE TABLE IF NOT EXISTS approval_matrix (
        id           INTEGER PRIMARY KEY AUTOINCREMENT,
        rating       TEXT NOT NULL,
        approver_role TEXT NOT NULL,
        approver_name TEXT,
        order_no     INTEGER DEFAULT 1,
        created_at   TEXT DEFAULT (datetime('now'))
    );

    CREATE TABLE IF NOT EXISTS contacts (
        id         INTEGER PRIMARY KEY AUTOINCREMENT,
        name       TEXT NOT NULL,
        email      TEXT NOT NULL,
        role       TEXT,
        active     INTEGER DEFAULT 1,
        alerts     TEXT DEFAULT '[]',
        created_at TEXT DEFAULT (datetime('now'))
    );

    CREATE TABLE IF NOT EXISTS notifications (
        id         INTEGER PRIMARY KEY AUTOINCREMENT,
        type       TEXT,
        message    TEXT,
        risk_no    TEXT,
        read       INTEGER DEFAULT 0,
        created_at TEXT DEFAULT (datetime('now'))
    );

    CREATE TABLE IF NOT EXISTS documents (
        id          INTEGER PRIMARY KEY AUTOINCREMENT,
        doc_type    TEXT NOT NULL,
        title       TEXT NOT NULL,
        filename    TEXT,
        content     TEXT,
        uploaded_by TEXT,
        created_at  TEXT DEFAULT (datetime('now')),
        updated_at  TEXT DEFAULT (datetime('now'))
    );

    CREATE TABLE IF NOT EXISTS impact_scale (
        id          INTEGER PRIMARY KEY AUTOINCREMENT,
        score       INTEGER NOT NULL UNIQUE,
        label       TEXT NOT NULL,
        description TEXT,
        fin_threshold TEXT
    );

    CREATE TABLE IF NOT EXISTS likelihood_scale (
        id          INTEGER PRIMARY KEY AUTOINCREMENT,
        score       INTEGER NOT NULL UNIQUE,
        label       TEXT NOT NULL,
        description TEXT,
        probability TEXT
    );
    """)

    # ── Idempotent column migrations (safe on existing DBs from older versions) ──
    def _ensure_col(table, col, ddl):
        cols = [r[1] for r in c.execute(f"PRAGMA table_info({table})").fetchall()]
        if col not in cols:
            c.execute(f"ALTER TABLE {table} ADD COLUMN {col} {ddl}")

    _ensure_col("risks",        "velocity",       "INTEGER")
    _ensure_col("risks",        "velocity_basis", "TEXT")
    _ensure_col("risks",        "risk_score_3d",  "INTEGER DEFAULT 0")
    _ensure_col("risks",        "rating_3d",      "TEXT DEFAULT 'N/A'")
    _ensure_col("risks",        "pending_status",                "TEXT")
    _ensure_col("risks",        "pending_status_submitted_by",   "TEXT")
    _ensure_col("risks",        "pending_status_submitted_at",   "TEXT")
    _ensure_col("risks",        "pending_status_approval",       "TEXT")
    _ensure_col("risk_history", "velocity",       "INTEGER")
    _ensure_col("risk_history", "score_3d",       "INTEGER")
    _ensure_col("risk_history", "rating_3d",      "TEXT")

    # Default settings
    defaults = {
        "company_name": "My Company",
        "company_logo": "",
        "report_period": "Q2 FY2025",
        "prepared_by": "Group Risk Management",
        "classification": "Strictly Confidential",
        "smtp_host": "", "smtp_port": "587",
        "smtp_user": "", "smtp_pass": "", "smtp_from": "",
        "scoring_mode_default": "2D",
    }
    for k, v in defaults.items():
        c.execute("INSERT OR IGNORE INTO settings (key,value) VALUES (?,?)", (k, v))

    # Default businesses
    for b in ["Group HQ","FinCo Capital","FinCo Insurance","FinCo Asset Mgmt","FinCo Retail","FinCo Tech"]:
        c.execute("INSERT OR IGNORE INTO businesses (name) VALUES (?)", (b,))

    # Default impact scale
    impact_defaults = [
        (1, "Negligible", "Minimal impact on operations or finances", "<$100K loss; no disruption"),
        (2, "Minor",      "Limited impact, quickly recoverable",       "$100K–$1M; short disruption"),
        (3, "Moderate",   "Significant impact, recovery possible",     "$1M–$10M; regulatory scrutiny"),
        (4, "Major",      "Severe impact, extended recovery needed",   "$10M–$50M; media coverage"),
        (5, "Catastrophic","Existential threat to the organisation",   ">$50M; existential threat"),
    ]
    for row in impact_defaults:
        c.execute("INSERT OR IGNORE INTO impact_scale (score,label,description,fin_threshold) VALUES (?,?,?,?)", row)

    # Default likelihood scale
    likelihood_defaults = [
        (1, "Rare",          "Highly unlikely to occur",            "<5% probability p.a."),
        (2, "Unlikely",      "Unlikely but could occur",            "5–25% probability"),
        (3, "Possible",      "May occur; has happened in industry", "25–50% probability"),
        (4, "Likely",        "Will probably occur; has occurred",   "50–75% probability"),
        (5, "Almost Certain","Expected to occur; recurring issue",  ">75% probability"),
    ]
    for row in likelihood_defaults:
        c.execute("INSERT OR IGNORE INTO likelihood_scale (score,label,description,probability) VALUES (?,?,?,?)", row)

    # Default velocity scale (per ERM 3D framework — time-to-impact)
    velocity_defaults = [
        (3, "High",   "Rapid Onset, little or no warning — Onset occurs in a matter of hours to 30 days from event to impact"),
        (2, "Medium", "Onset occurs in a matter of 30 days to 90 days from event to impact"),
        (1, "Low",    "Onset occurs in a matter of 90 days to a year from event to impact"),
    ]
    for row in velocity_defaults:
        c.execute("INSERT OR IGNORE INTO velocity_scale (score,label,description) VALUES (?,?,?)", row)

    # Default approval matrix
    matrix_defaults = [
        ("Critical", "Board Risk Committee", "Board Chairman", 1),
        ("Critical", "Group CRO", "Chief Risk Officer", 2),
        ("High",     "Group CRO", "Chief Risk Officer", 1),
        ("High",     "Business Head", "", 2),
        ("Medium",   "Risk Officer", "", 1),
        ("Low",      "Risk Officer", "", 1),
    ]
    for row in matrix_defaults:
        c.execute("INSERT OR IGNORE INTO approval_matrix (rating,approver_role,approver_name,order_no) VALUES (?,?,?,?)", row)

    conn.commit()
    conn.close()

# ── Helpers ──────────────────────────────────────────────────────────────────

def _rating_2d(score):
    """2D bands (Impact × Likelihood, max=25)."""
    if score >= 21: return "Critical"
    if score >= 15: return "High"
    if score >= 10:  return "Medium"
    if score >= 5:  return "Low"
    return "N/A"

def _rating_3d(score):
    """3D bands (Impact × Likelihood × Velocity, max=75) — per ERM framework."""
    if score >= 51: return "Critical"
    if score >= 31: return "High"
    if score >= 16: return "Medium"
    if score >= 1:  return "Low"
    return "N/A"

def compute_score(impacts, likelihood, velocity=None):
    """
    Returns (final_impact, score_2d, rating_2d, score_3d, rating_3d, velocity_int).
    score_3d/rating_3d are 0/'N/A' if velocity is missing — keeps platform back-compat.
    """
    vals = []
    for v in impacts.values():
        if v and v != "NA":
            try: vals.append(int(v))
            except: pass
    fi = max(vals) if vals else 0
    l  = int(likelihood) if likelihood else 0
    score_2d  = fi * l if fi and l else 0
    rating_2d = _rating_2d(score_2d)
    v = None
    try:
        v = int(velocity) if velocity not in (None, "", "NA") else None
    except: v = None
    score_3d  = (score_2d * v) if (score_2d and v) else 0
    rating_3d = _rating_3d(score_3d) if score_3d else "N/A"
    return fi, score_2d, rating_2d, score_3d, rating_3d, v

def row_to_dict(row):
    if row is None: return None
    d = dict(row)
    for k in ["businesses","alerts","attachments"]:
        if k in d and isinstance(d[k], str):
            try: d[k] = json.loads(d[k])
            except: d[k] = []
    return d

def add_notification(conn, ntype, message, risk_no=""):
    conn.execute("INSERT INTO notifications (type,message,risk_no) VALUES (?,?,?)", (ntype, message, risk_no))

# ── Risk CRUD ─────────────────────────────────────────────────────────────────

def get_risks():
    conn = get_db()
    risks = [row_to_dict(r) for r in conn.execute("SELECT * FROM risks ORDER BY no").fetchall()]
    for r in risks:
        r["history"] = [row_to_dict(h) for h in conn.execute(
            "SELECT * FROM risk_history WHERE risk_id=? ORDER BY date", (r["id"],)).fetchall()]
        r["approvals"] = [row_to_dict(a) for a in conn.execute(
            "SELECT * FROM approvals WHERE risk_id=? ORDER BY created_at DESC", (r["id"],)).fetchall()]
        r["reviews"] = [row_to_dict(rv) for rv in conn.execute(
            "SELECT * FROM review_periods WHERE risk_id=? ORDER BY created_at", (r["id"],)).fetchall()]
    conn.close()
    return risks

def get_risk(rid):
    conn = get_db()
    r = row_to_dict(conn.execute("SELECT * FROM risks WHERE id=?", (rid,)).fetchone())
    if r:
        r["history"]   = [row_to_dict(h) for h in conn.execute("SELECT * FROM risk_history WHERE risk_id=? ORDER BY date", (rid,)).fetchall()]
        r["approvals"] = [row_to_dict(a) for a in conn.execute("SELECT * FROM approvals WHERE risk_id=? ORDER BY created_at DESC", (rid,)).fetchall()]
        r["reviews"]   = [row_to_dict(rv) for rv in conn.execute("SELECT * FROM review_periods WHERE risk_id=? ORDER BY created_at", (rid,)).fetchall()]
    conn.close()
    return r

def create_risk(data: dict):
    conn = get_db()
    impacts = {
        "Financial":       data.get("impact_financial",""),
        "Operational":     data.get("impact_operational",""),
        "Regulatory":      data.get("impact_regulatory",""),
        "Brand":           data.get("impact_brand",""),
    }
    fi, s2, r2, s3, r3, v = compute_score(impacts, data.get("likelihood"), data.get("velocity"))
    now = datetime.now().strftime("%Y-%m")
    cur = conn.execute("""
        INSERT INTO risks (no,category,statement,events,
            impact_financial,impact_financial_basis,
            impact_operational,impact_operational_basis,
            impact_regulatory,impact_regulatory_basis,
            impact_brand,impact_brand_basis,
            final_impact,likelihood,likelihood_basis,
            velocity,velocity_basis,
            risk_score,rating,risk_score_3d,rating_3d,
            applicability,businesses,risk_owner,strategy,existing_controls,
            proposed_plan,proposed_plan_submitted_by,proposed_plan_submitted_at,
            mitigation_owner,target_date,status,actual_date,
            assessment_frequency,next_assessment_date,added_by)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
    """, (
        data.get("no"), data.get("category"), data.get("statement"), data.get("events"),
        data.get("impact_financial"), data.get("impact_financial_basis"),
        data.get("impact_operational"), data.get("impact_operational_basis"),
        data.get("impact_regulatory"), data.get("impact_regulatory_basis"),
        data.get("impact_brand"), data.get("impact_brand_basis"),
        fi, data.get("likelihood"), data.get("likelihood_basis"),
        v, data.get("velocity_basis"),
        s2, r2, s3, r3,
        data.get("applicability","group"),
        json.dumps(data.get("businesses",[])),
        data.get("risk_owner"), data.get("strategy"), data.get("existing_controls"),
        data.get("proposed_plan"),
        data.get("proposed_plan_submitted_by"), data.get("proposed_plan_submitted_at"),
        data.get("mitigation_owner"), data.get("target_date"),
        data.get("status","Not Started"), data.get("actual_date"),
        data.get("assessment_frequency","Quarterly"), data.get("next_assessment_date"),
        data.get("added_by",""),
    ))
    rid = cur.lastrowid
    conn.execute("""INSERT INTO risk_history
        (risk_id,date,impact,likelihood,velocity,score,rating,score_3d,rating_3d,note,changed_by)
        VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
        (rid, now, fi, data.get("likelihood"), v, s2, r2, s3, r3, "Initial assessment", data.get("added_by","")))
    add_notification(conn, "New Risk Added", f"New risk {data.get('no')} added: {(data.get('statement') or '')[:60]}", data.get("no",""))
    conn.commit()
    rid2 = rid
    conn.close()
    return get_risk(rid2)

def update_risk(rid, data: dict):
    conn = get_db()
    impacts = {
        "Financial":   data.get("impact_financial",""),
        "Operational": data.get("impact_operational",""),
        "Regulatory":  data.get("impact_regulatory",""),
        "Brand":       data.get("impact_brand",""),
    }
    fi, s2, r2, s3, r3, v = compute_score(impacts, data.get("likelihood"), data.get("velocity"))
    now = datetime.now().strftime("%Y-%m")
    conn.execute("""
        UPDATE risks SET no=?,category=?,statement=?,events=?,
            impact_financial=?,impact_financial_basis=?,
            impact_operational=?,impact_operational_basis=?,
            impact_regulatory=?,impact_regulatory_basis=?,
            impact_brand=?,impact_brand_basis=?,
            final_impact=?,likelihood=?,likelihood_basis=?,
            velocity=?,velocity_basis=?,
            risk_score=?,rating=?,risk_score_3d=?,rating_3d=?,
            applicability=?,businesses=?,risk_owner=?,strategy=?,existing_controls=?,
            proposed_plan=?,proposed_plan_submitted_by=?,proposed_plan_submitted_at=?,
            mitigation_owner=?,target_date=?,status=?,actual_date=?,
            assessment_frequency=?,next_assessment_date=?,
            updated_at=datetime('now')
        WHERE id=?
    """, (
        data.get("no"), data.get("category"), data.get("statement"), data.get("events"),
        data.get("impact_financial"), data.get("impact_financial_basis"),
        data.get("impact_operational"), data.get("impact_operational_basis"),
        data.get("impact_regulatory"), data.get("impact_regulatory_basis"),
        data.get("impact_brand"), data.get("impact_brand_basis"),
        fi, data.get("likelihood"), data.get("likelihood_basis"),
        v, data.get("velocity_basis"),
        s2, r2, s3, r3,
        data.get("applicability","group"),
        json.dumps(data.get("businesses",[])),
        data.get("risk_owner"), data.get("strategy"), data.get("existing_controls"),
        data.get("proposed_plan"),
        data.get("proposed_plan_submitted_by"), data.get("proposed_plan_submitted_at"),
        data.get("mitigation_owner"), data.get("target_date"),
        data.get("status"), data.get("actual_date"),
        data.get("assessment_frequency","Quarterly"), data.get("next_assessment_date"),
        rid,
    ))
    conn.execute("""INSERT INTO risk_history
        (risk_id,date,impact,likelihood,velocity,score,rating,score_3d,rating_3d,note,changed_by)
        VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
        (rid, now, fi, data.get("likelihood"), v, s2, r2, s3, r3,
         data.get("history_note","Updated"), data.get("changed_by","")))
    add_notification(conn, "Risk Score Changed", f"Risk {data.get('no')} updated — score: {s2} ({r2})", data.get("no",""))
    conn.commit()
    conn.close()
    return get_risk(rid)

def delete_risk(rid):
    conn = get_db()
    conn.execute("DELETE FROM risk_history WHERE risk_id=?", (rid,))
    conn.execute("DELETE FROM approvals WHERE risk_id=?", (rid,))
    conn.execute("DELETE FROM review_periods WHERE risk_id=?", (rid,))
    conn.execute("DELETE FROM risks WHERE id=?", (rid,))
    conn.commit()
    conn.close()

# ── Approvals ─────────────────────────────────────────────────────────────────

def get_approval_matrix():
    conn = get_db()
    rows = [row_to_dict(r) for r in conn.execute("SELECT * FROM approval_matrix ORDER BY rating,order_no").fetchall()]
    conn.close()
    return rows

def save_approval_matrix(rows):
    conn = get_db()
    conn.execute("DELETE FROM approval_matrix")
    for r in rows:
        conn.execute("INSERT INTO approval_matrix (rating,approver_role,approver_name,order_no) VALUES (?,?,?,?)",
            (r["rating"], r["approver_role"], r.get("approver_name",""), r.get("order_no",1)))
    conn.commit()
    conn.close()

def approve_risk(rid, data: dict):
    conn = get_db()
    conn.execute("""
        INSERT INTO approvals (risk_id,status,approver,approved_by,approved_at,comments)
        VALUES (?,?,?,?,datetime('now'),?)
    """, (rid, data["status"], data["approver"], data.get("approved_by",""), data.get("comments","")))
    if data["status"] == "Approved":
        conn.execute("UPDATE risks SET status='Completed',actual_date=date('now'),updated_at=datetime('now') WHERE id=?", (rid,))
        add_notification(conn, "Rating Escalated", f"Risk closure approved by {data['approver']}", "")
    elif data["status"] == "Rejected":
        conn.execute("UPDATE risks SET status='In Progress',updated_at=datetime('now') WHERE id=?", (rid,))
    conn.commit()
    conn.close()

# ── Settings ──────────────────────────────────────────────────────────────────

def get_settings():
    conn = get_db()
    rows = conn.execute("SELECT * FROM settings").fetchall()
    conn.close()
    return {r["key"]: r["value"] for r in rows}

def save_settings(data: dict):
    conn = get_db()
    for k, v in data.items():
        conn.execute("INSERT OR REPLACE INTO settings (key,value) VALUES (?,?)", (k, str(v)))
    conn.commit()
    conn.close()

# ── Businesses ────────────────────────────────────────────────────────────────

def get_businesses():
    conn = get_db()
    rows = [row_to_dict(r) for r in conn.execute("SELECT * FROM businesses ORDER BY name").fetchall()]
    conn.close()
    return rows

def create_business(name):
    conn = get_db()
    conn.execute("INSERT OR IGNORE INTO businesses (name) VALUES (?)", (name,))
    conn.commit()
    conn.close()

def update_business(bid, name, active):
    conn = get_db()
    conn.execute("UPDATE businesses SET name=?,active=? WHERE id=?", (name, active, bid))
    conn.commit()
    conn.close()

def delete_business(bid):
    conn = get_db()
    conn.execute("DELETE FROM businesses WHERE id=?", (bid,))
    conn.commit()
    conn.close()

# ── Scales ────────────────────────────────────────────────────────────────────

def get_impact_scale():
    conn = get_db()
    rows = [row_to_dict(r) for r in conn.execute("SELECT * FROM impact_scale ORDER BY score").fetchall()]
    conn.close()
    return rows

def get_likelihood_scale():
    conn = get_db()
    rows = [row_to_dict(r) for r in conn.execute("SELECT * FROM likelihood_scale ORDER BY score").fetchall()]
    conn.close()
    return rows

def save_impact_scale(rows):
    conn = get_db()
    conn.execute("DELETE FROM impact_scale")
    for r in rows:
        conn.execute("INSERT INTO impact_scale (score,label,description,fin_threshold) VALUES (?,?,?,?)",
            (r["score"], r["label"], r.get("description",""), r.get("fin_threshold","")))
    conn.commit()
    conn.close()

def save_likelihood_scale(rows):
    conn = get_db()
    conn.execute("DELETE FROM likelihood_scale")
    for r in rows:
        conn.execute("INSERT INTO likelihood_scale (score,label,description,probability) VALUES (?,?,?,?)",
            (r["score"], r["label"], r.get("description",""), r.get("probability","")))
    conn.commit()
    conn.close()

def get_velocity_scale():
    conn = get_db()
    rows = [row_to_dict(r) for r in conn.execute("SELECT * FROM velocity_scale ORDER BY score DESC").fetchall()]
    conn.close()
    return rows

def save_velocity_scale(rows):
    conn = get_db()
    conn.execute("DELETE FROM velocity_scale")
    for r in rows:
        conn.execute("INSERT INTO velocity_scale (score,label,description) VALUES (?,?,?)",
            (r["score"], r["label"], r.get("description","")))
    conn.commit()
    conn.close()

# ── Review Periods (periodic review workflow) ────────────────────────────────

def _next_due_date(current_due, frequency):
    """Advance next_assessment_date based on frequency."""
    from datetime import date, timedelta
    base = None
    if current_due:
        try: base = datetime.strptime(current_due, "%Y-%m-%d").date()
        except: base = None
    if not base: base = date.today()
    days = {"Monthly":30, "Quarterly":91, "Semi-Annually":182, "Annually":365}.get(frequency, 91)
    return (base + timedelta(days=days)).isoformat()

def _period_label(due_date, frequency):
    """Generate a period label e.g. 'FY26-Q1' (Indian FY: Apr–Mar) or '2026-05'."""
    from datetime import date
    try:
        d = datetime.strptime(due_date, "%Y-%m-%d").date() if due_date else date.today()
    except:
        d = date.today()
    if frequency == "Monthly":
        return d.strftime("%Y-%m")
    if frequency in ("Quarterly","Semi-Annually"):
        if d.month >= 4:
            fy = d.year + 1
        else:
            fy = d.year
        q_idx = ((d.month - 4) % 12) // 3 + 1
        if frequency == "Semi-Annually":
            half = 1 if q_idx <= 2 else 2
            return f"FY{str(fy)[-2:]}-H{half}"
        return f"FY{str(fy)[-2:]}-Q{q_idx}"
    return f"FY{str(d.year + (1 if d.month>=4 else 0))[-2:]}"

def get_pending_reviews():
    """Reviews awaiting submission (Due) or awaiting approver decision (Submitted)."""
    conn = get_db()
    rows = [row_to_dict(r) for r in conn.execute("""
        SELECT rp.*, r.no AS risk_no, r.statement AS risk_statement, r.risk_owner, r.rating AS risk_rating
        FROM review_periods rp JOIN risks r ON rp.risk_id = r.id
        WHERE rp.status IN ('Due','Submitted')
        ORDER BY rp.due_date, rp.created_at
    """).fetchall()]
    conn.close()
    return rows

def get_risk_reviews(rid):
    conn = get_db()
    rows = [row_to_dict(r) for r in conn.execute(
        "SELECT * FROM review_periods WHERE risk_id=? ORDER BY created_at", (rid,)).fetchall()]
    conn.close()
    return rows

def open_review_if_due(rid):
    """Open a 'Due' review row if next_assessment_date has arrived and none open."""
    from datetime import date
    conn = get_db()
    risk = conn.execute("SELECT * FROM risks WHERE id=?", (rid,)).fetchone()
    if not risk:
        conn.close(); return None
    due = risk["next_assessment_date"]
    if not due:
        conn.close(); return None
    try:
        due_d = datetime.strptime(due, "%Y-%m-%d").date()
    except:
        conn.close(); return None
    if due_d > date.today():
        conn.close(); return None
    existing = conn.execute(
        "SELECT id FROM review_periods WHERE risk_id=? AND status IN ('Due','Submitted')", (rid,)
    ).fetchone()
    if existing:
        conn.close(); return existing[0]
    label = _period_label(due, risk["assessment_frequency"] or "Quarterly")
    cur = conn.execute("INSERT INTO review_periods (risk_id,period_label,due_date,status) VALUES (?,?,?,?)",
                       (rid, label, due, "Due"))
    add_notification(conn, "Assessment Due",
                     f"Review due for risk {risk['no']} — period {label}", risk["no"])
    conn.commit()
    rvid = cur.lastrowid
    conn.close()
    return rvid

def open_due_reviews_all():
    """Sweep all risks; open Due rows for any past their next_assessment_date."""
    from datetime import date
    conn = get_db()
    today = date.today().isoformat()
    # Find risks that are past due and don't already have an open review
    risks = conn.execute("""
        SELECT r.id, r.no, r.next_assessment_date, r.assessment_frequency
        FROM risks r
        WHERE r.next_assessment_date IS NOT NULL
          AND r.next_assessment_date <= ?
          AND NOT EXISTS (
              SELECT 1 FROM review_periods rp
              WHERE rp.risk_id = r.id AND rp.status IN ('Due','Submitted')
          )
    """, (today,)).fetchall()
    opened = 0
    for risk in risks:
        label = _period_label(risk["next_assessment_date"], risk["assessment_frequency"] or "Quarterly")
        conn.execute("INSERT INTO review_periods (risk_id,period_label,due_date,status) VALUES (?,?,?,?)",
                     (risk["id"], label, risk["next_assessment_date"], "Due"))
        add_notification(conn, "Assessment Due",
                         f"Review due for risk {risk['no']} — period {label}", risk["no"])
        opened += 1
    if opened:
        conn.commit()
    conn.close()
    return opened

def submit_review(rid, data):
    """Owner submits a review for the current open period (or creates one)."""
    conn = get_db()
    risk = row_to_dict(conn.execute("SELECT * FROM risks WHERE id=?", (rid,)).fetchone())
    if not risk:
        conn.close(); return None
    open_row = conn.execute(
        "SELECT id,period_label,due_date FROM review_periods WHERE risk_id=? AND status IN ('Due','Submitted') ORDER BY created_at DESC LIMIT 1",
        (rid,)
    ).fetchone()
    impacts = {
        "Financial":   data.get("impact_financial",""),
        "Operational": data.get("impact_operational",""),
        "Regulatory":  data.get("impact_regulatory",""),
        "Brand":       data.get("impact_brand",""),
    }
    fi, s2, r2, s3, r3, v = compute_score(impacts, data.get("likelihood"), data.get("velocity"))
    atts = data.get("attachments", [])
    if open_row:
        rvid = open_row["id"]
        period = open_row["period_label"]
        conn.execute("""UPDATE review_periods SET status='Submitted',
            submitted_by=?, submitted_at=datetime('now'),
            final_impact=?, impact_financial=?, impact_operational=?, impact_regulatory=?, impact_brand=?,
            likelihood=?, velocity=?, score_2d=?, rating_2d=?, score_3d=?, rating_3d=?,
            comments=?, attachments=?
            WHERE id=?""",
            (data.get("submitted_by",""), fi,
             data.get("impact_financial"), data.get("impact_operational"),
             data.get("impact_regulatory"), data.get("impact_brand"),
             data.get("likelihood"), v, s2, r2, s3, r3,
             data.get("comments",""), json.dumps(atts), rvid))
    else:
        period = _period_label(risk.get("next_assessment_date") or "", risk.get("assessment_frequency") or "Quarterly")
        due    = risk.get("next_assessment_date")
        cur = conn.execute("""INSERT INTO review_periods
            (risk_id,period_label,due_date,status,submitted_by,submitted_at,
             final_impact,impact_financial,impact_operational,impact_regulatory,impact_brand,
             likelihood,velocity,score_2d,rating_2d,score_3d,rating_3d,comments,attachments)
            VALUES (?,?,?,?,?,datetime('now'),?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (rid, period, due, "Submitted", data.get("submitted_by",""), fi,
             data.get("impact_financial"), data.get("impact_operational"),
             data.get("impact_regulatory"), data.get("impact_brand"),
             data.get("likelihood"), v, s2, r2, s3, r3, data.get("comments",""), json.dumps(atts)))
        rvid = cur.lastrowid
    add_notification(conn, "Review Submitted",
                     f"Review submitted for {risk['no']} ({period}) — pending approval", risk["no"])
    conn.commit()
    conn.close()
    return rvid

def decide_review(review_id, decision, approver_role, approved_by, comments):
    """Approver Approves/Rejects a submitted review.
       On Approve: rating becomes the source-of-truth for the period; risk's primary
       fields are updated; next_assessment_date advances; risk_history gets a row.
       On Reject: a fresh Due period is opened so owner can resubmit.
    """
    conn = get_db()
    rv = row_to_dict(conn.execute("SELECT * FROM review_periods WHERE id=?", (review_id,)).fetchone())
    if not rv:
        conn.close(); return None
    risk = row_to_dict(conn.execute("SELECT * FROM risks WHERE id=?", (rv["risk_id"],)).fetchone())
    now_ts = datetime.now().strftime("%Y-%m")
    if decision == "Approved":
        conn.execute("""UPDATE review_periods SET status='Approved',
            approver_role=?, approved_by=?, decided_at=datetime('now'), decision_comments=? WHERE id=?""",
            (approver_role or "", approved_by or "", comments or "", review_id))
        new_due = _next_due_date(rv.get("due_date"), risk.get("assessment_frequency") or "Quarterly")
        conn.execute("""UPDATE risks SET
            impact_financial=COALESCE(?,impact_financial),
            impact_operational=COALESCE(?,impact_operational),
            impact_regulatory=COALESCE(?,impact_regulatory),
            impact_brand=COALESCE(?,impact_brand),
            final_impact=?, likelihood=?, velocity=?,
            risk_score=?, rating=?, risk_score_3d=?, rating_3d=?,
            next_assessment_date=?, updated_at=datetime('now')
            WHERE id=?""",
            (rv.get("impact_financial"), rv.get("impact_operational"),
             rv.get("impact_regulatory"), rv.get("impact_brand"),
             rv.get("final_impact"), rv.get("likelihood"), rv.get("velocity"),
             rv.get("score_2d"), rv.get("rating_2d"),
             rv.get("score_3d"), rv.get("rating_3d"),
             new_due, rv["risk_id"]))
        conn.execute("""INSERT INTO risk_history
            (risk_id,date,impact,likelihood,velocity,score,rating,score_3d,rating_3d,note,changed_by)
            VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
            (rv["risk_id"], now_ts, rv.get("final_impact"), rv.get("likelihood"), rv.get("velocity"),
             rv.get("score_2d"), rv.get("rating_2d"), rv.get("score_3d"), rv.get("rating_3d"),
             f"Period review approved ({rv['period_label']})", approved_by or ""))
        add_notification(conn, "Rating Escalated",
                         f"Review approved for {risk['no']} ({rv['period_label']}) — score {rv.get('score_2d')} ({rv.get('rating_2d')})",
                         risk["no"])
    elif decision == "Rejected":
        conn.execute("""UPDATE review_periods SET status='Rejected',
            approver_role=?, approved_by=?, decided_at=datetime('now'), decision_comments=? WHERE id=?""",
            (approver_role or "", approved_by or "", comments or "", review_id))
        conn.execute("INSERT INTO review_periods (risk_id,period_label,due_date,status) VALUES (?,?,?,?)",
                     (rv["risk_id"], rv["period_label"], rv.get("due_date"), "Due"))
        add_notification(conn, "Review Rejected",
                         f"Review for {risk['no']} ({rv['period_label']}) rejected — please resubmit", risk["no"])
    conn.commit()
    conn.close()
    return True

def get_portfolio_movement():
    """Compare each risk's last two Approved reviews (or 1 Approved vs initial assessment).
       Returns counts {up, down, same, new} and per-risk detail.
    """
    conn = get_db()
    risks = [row_to_dict(r) for r in conn.execute("SELECT id,no,statement,risk_score,rating FROM risks").fetchall()]
    detail = []
    up = down = same = new_ = 0
    for r in risks:
        approved = [row_to_dict(x) for x in conn.execute(
            "SELECT score_2d,rating_2d,score_3d,rating_3d,period_label,decided_at FROM review_periods "
            "WHERE risk_id=? AND status='Approved' ORDER BY decided_at", (r["id"],)).fetchall()]
        if len(approved) == 0:
            new_ += 1
            detail.append({**r, "movement":"new", "delta":0,
                           "prev_score":None, "curr_score":r["risk_score"],
                           "prev_period":None, "curr_period":"Initial"})
            continue
        if len(approved) == 1:
            init = conn.execute(
                "SELECT score FROM risk_history WHERE risk_id=? ORDER BY id LIMIT 1", (r["id"],)
            ).fetchone()
            prev_score = init["score"] if init else None
            curr_score = approved[0]["score_2d"]
            prev_period = "Initial"
        else:
            prev_score = approved[-2]["score_2d"]
            curr_score = approved[-1]["score_2d"]
            prev_period = approved[-2]["period_label"]
        delta = (curr_score or 0) - (prev_score or 0)
        if delta > 0: up += 1; mv = "up"
        elif delta < 0: down += 1; mv = "down"
        else: same += 1; mv = "same"
        detail.append({**r, "movement":mv, "delta":delta,
                       "prev_score":prev_score, "curr_score":curr_score,
                       "prev_period":prev_period, "curr_period":approved[-1]["period_label"]})
    conn.close()
    return {"up":up, "down":down, "same":same, "new":new_, "detail":detail}

# ── Contacts ──────────────────────────────────────────────────────────────────

def get_contacts():
    conn = get_db()
    rows = [row_to_dict(r) for r in conn.execute("SELECT * FROM contacts ORDER BY id").fetchall()]
    conn.close()
    return rows

def upsert_contact(data):
    conn = get_db()
    if data.get("id"):
        conn.execute("UPDATE contacts SET name=?,email=?,role=?,active=?,alerts=? WHERE id=?",
            (data["name"],data["email"],data.get("role",""),data.get("active",1),
             json.dumps(data.get("alerts",[])),data["id"]))
    else:
        conn.execute("INSERT INTO contacts (name,email,role,active,alerts) VALUES (?,?,?,?,?)",
            (data["name"],data["email"],data.get("role",""),data.get("active",1),
             json.dumps(data.get("alerts",[]))))
    conn.commit()
    conn.close()

def delete_contact(cid):
    conn = get_db()
    conn.execute("DELETE FROM contacts WHERE id=?", (cid,))
    conn.commit()
    conn.close()

# ── Notifications ─────────────────────────────────────────────────────────────

def get_notifications():
    conn = get_db()
    rows = [row_to_dict(r) for r in conn.execute("SELECT * FROM notifications ORDER BY created_at DESC LIMIT 100").fetchall()]
    conn.close()
    return rows

def clear_notifications():
    conn = get_db()
    conn.execute("DELETE FROM notifications")
    conn.commit()
    conn.close()

# ── Documents ─────────────────────────────────────────────────────────────────

def get_documents():
    conn = get_db()
    rows = [row_to_dict(r) for r in conn.execute("SELECT * FROM documents ORDER BY doc_type,title").fetchall()]
    conn.close()
    return rows

def upsert_document(data):
    conn = get_db()
    if data.get("id"):
        conn.execute("UPDATE documents SET doc_type=?,title=?,content=?,filename=?,updated_at=datetime('now') WHERE id=?",
            (data["doc_type"],data["title"],data.get("content",""),data.get("filename",""),data["id"]))
    else:
        conn.execute("INSERT INTO documents (doc_type,title,content,filename,uploaded_by) VALUES (?,?,?,?,?)",
            (data["doc_type"],data["title"],data.get("content",""),data.get("filename",""),data.get("uploaded_by","")))
    conn.commit()
    conn.close()

def delete_document(did):
    conn = get_db()
    conn.execute("DELETE FROM documents WHERE id=?", (did,))
    conn.commit()
    conn.close()

# ── Excel Import ──────────────────────────────────────────────────────────────

def import_risks_from_excel(filepath, added_by=""):
    from openpyxl import load_workbook
    wb = load_workbook(filepath, data_only=True)
    ws = wb.active
    headers = [str(c.value).strip() if c.value else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
    imported, errors = 0, []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        row_data = {headers[j]: (row[j] if j < len(row) else None) for j in range(len(headers))}
        if not row_data.get("no") and not row_data.get("No") and not row_data.get("Risk No"):
            continue
        try:
            data = {
                "no":                       str(row_data.get("no") or row_data.get("No") or row_data.get("Risk No") or ""),
                "category":                 str(row_data.get("category") or row_data.get("Category") or ""),
                "statement":                str(row_data.get("statement") or row_data.get("Risk Statement") or ""),
                "events":                   str(row_data.get("events") or row_data.get("Risk Events") or ""),
                "impact_financial":         str(row_data.get("impact_financial") or row_data.get("Financial Impact") or ""),
                "impact_financial_basis":   str(row_data.get("impact_financial_basis") or ""),
                "impact_operational":       str(row_data.get("impact_operational") or row_data.get("Operational Impact") or ""),
                "impact_operational_basis": str(row_data.get("impact_operational_basis") or ""),
                "impact_regulatory":        str(row_data.get("impact_regulatory") or row_data.get("Regulatory Impact") or ""),
                "impact_regulatory_basis":  str(row_data.get("impact_regulatory_basis") or ""),
                "impact_brand":             str(row_data.get("impact_brand") or row_data.get("Brand Impact") or ""),
                "impact_brand_basis":       str(row_data.get("impact_brand_basis") or ""),
                "likelihood":               row_data.get("likelihood") or row_data.get("Likelihood"),
                "likelihood_basis":         str(row_data.get("likelihood_basis") or row_data.get("Likelihood Basis") or ""),
                "velocity":                 row_data.get("velocity") or row_data.get("Velocity"),
                "velocity_basis":           str(row_data.get("velocity_basis") or row_data.get("Velocity Basis") or ""),
                "applicability":            str(row_data.get("applicability") or "group"),
                "businesses":              [x.strip() for x in str(row_data.get("businesses","")).split(";") if x.strip()],
                "risk_owner":               str(row_data.get("risk_owner") or row_data.get("Risk Owner") or ""),
                "strategy":                 str(row_data.get("strategy") or row_data.get("Strategy") or ""),
                "existing_controls":        str(row_data.get("existing_controls") or row_data.get("Existing Controls") or ""),
                "proposed_plan":            str(row_data.get("proposed_plan") or row_data.get("Proposed Plan") or ""),
                "mitigation_owner":         str(row_data.get("mitigation_owner") or row_data.get("Mitigation Owner") or ""),
                "target_date":              str(row_data.get("target_date") or row_data.get("Target Date") or ""),
                "status":                   str(row_data.get("status") or row_data.get("Status") or "Not Started"),
                "assessment_frequency":     str(row_data.get("assessment_frequency") or "Quarterly"),
                "next_assessment_date":     str(row_data.get("next_assessment_date") or ""),
                "added_by":                 added_by,
            }
            create_risk(data)
            imported += 1
        except Exception as e:
            errors.append(f"Row {i}: {str(e)}")
    return imported, errors

def export_risks_csv():
    import csv, io
    risks = get_risks()
    out = io.StringIO()
    headers = [
        "no","category","statement","events",
        "impact_financial","impact_financial_basis",
        "impact_operational","impact_operational_basis",
        "impact_regulatory","impact_regulatory_basis",
        "impact_brand","impact_brand_basis",
        "final_impact","likelihood","likelihood_basis",
        "velocity","velocity_basis",
        "risk_score","rating","risk_score_3d","rating_3d",
        "applicability","businesses",
        "risk_owner","strategy","existing_controls","proposed_plan",
        "mitigation_owner","target_date","status","actual_date",
        "assessment_frequency","next_assessment_date","added_by","created_at",
    ]
    w = csv.DictWriter(out, fieldnames=headers, extrasaction="ignore")
    w.writeheader()
    for r in risks:
        r2 = dict(r)
        r2["businesses"] = "; ".join(r2.get("businesses") or [])
        w.writerow({h: r2.get(h,"") for h in headers})
    return out.getvalue()

def generate_excel_template():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    import io
    wb = Workbook()
    ws = wb.active
    ws.title = "Risk Register"
    headers = [
        "no","category","statement","events",
        "impact_financial","impact_financial_basis",
        "impact_operational","impact_operational_basis",
        "impact_regulatory","impact_regulatory_basis",
        "impact_brand","impact_brand_basis",
        "likelihood","likelihood_basis",
        "velocity","velocity_basis",
        "applicability","businesses","risk_owner","strategy",
        "existing_controls","proposed_plan","mitigation_owner",
        "target_date","status","assessment_frequency","next_assessment_date",
    ]
    hfill  = PatternFill("solid", fgColor="0D2340")
    hfont  = Font(color="FFFFFF", bold=True, size=10)
    afill  = PatternFill("solid", fgColor="FDB913")
    afont  = Font(color="0D2340", bold=True, size=10)
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = hfill
        cell.font = hfont
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = max(len(h)+4, 18)
    # Sample row
    sample = ["R-001","Technology","Cybersecurity breach","Ransomware, phishing",
              "5","High financial exposure","5","Critical ops dependency",
              "4","Regulatory fines likely","5","Reputational damage severe",
              "3","Near-miss incidents in FY24",
              "3","Rapid onset; ransomware unfolds within hours",
              "group","","Group CISO","Mitigate",
              "EDR, SOC, MFA","Zero-trust architecture","Group CISO",
              "2025-12-31","Not Started","Quarterly","2025-06-30"]
    for col, val in enumerate(sample, 1):
        cell = ws.cell(row=2, column=col, value=val)
        cell.fill = afill
        cell.font = afont
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
